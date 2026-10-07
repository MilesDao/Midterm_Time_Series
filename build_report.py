"""Fill GroupProjectTemplate2026.docx with code/results from Group5_TimeSeries.ipynb -> Group5_Report.docx."""
import base64
import io

import nbformat
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from report_text import CONTENT, STUDENTS, TITLE

FONT = "Lucida Sans"
nb = nbformat.read("Group5_TimeSeries.ipynb", 4)
CELLS = {c.metadata.tags[0]: c for c in nb.cells if c.cell_type == "code"}

doc = Document("GroupProjectTemplate2026.docx")
body = doc.element.body

# ---- cover page: the template anchors the cover text box to a paragraph deep in the outline;
# keep only that paragraph's drawing runs, move it first, fill group number and title
cover = next(p._p for p in doc.paragraphs if "GROUP PROJECT REPORT" in p._p.xml)
for r in list(cover.iterchildren(qn("w:r"))):
    if not r.xpath(".//w:drawing"):
        cover.remove(r)
if cover.pPr is not None:
    cover.remove(cover.pPr)
body.insert(0, cover)
for p in list(cover.iter(qn("w:p"))):
    if p is cover or p.find(".//" + qn("w:p")) is not None:   # only innermost text-box paragraphs
        continue
    ts = list(p.iter(qn("w:t")))
    text = "".join(t.text for t in ts)
    new = None
    if "Student group number" in text:
        new = "Student group number 5"
    elif text.startswith("Time Series Analysis of"):
        new = TITLE
    elif text.startswith("Student ") and " name" in text:   # "Student N name ... – Student ID N"
        n = int(text.split()[1])
        if n > len(STUDENTS):
            p.getparent().remove(p)
            continue
        new = f"{STUDENTS[n - 1][0]} – Student ID {STUDENTS[n - 1][1]}"
    if new and ts:
        ts[0].text = new
        for t in ts[1:]:
            t.text = ""
    # keep the text box single-spaced (Normal becomes 1.5 below and would overflow the box)
    spacing = p.get_or_add_pPr().get_or_add_spacing()
    spacing.set(qn("w:line"), "240"); spacing.set(qn("w:lineRule"), "auto")
    for r in p.iter(qn("w:r")):   # cover keeps the template's Times look so the date still fits
        rf = r.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
    if text.startswith("Hanoi, October"):   # drop 2 spacer lines above the date so it stays inside the box
        for _ in range(2):
            prev = p.getprevious()
            if prev is not None and prev.tag == qn("w:p") and not "".join(prev.itertext()).strip():
                prev.getparent().remove(prev)
# logo: the template squeezes the 16:9 png into a 2.9:1 box -> restore its ratio, and move it to the top
for logo_p in cover.xpath(".//w:p[.//pic:pic][not(.//w:p)]"):
    for ext in logo_p.xpath(".//wp:extent | .//a:ext[@cx]"):
        ext.set("cy", str(round(int(ext.get("cx")) * 9 / 16)))
    univ = next(q for q in logo_p.getparent().iterchildren(qn("w:p"))
                if "".join(q.itertext()).startswith("UNIVERSITY OF SCIENCE"))
    univ.addprevious(logo_p)
for el in list(body):
    if el is not cover and el.tag != qn("w:sectPr"):
        body.remove(el)

# ---- styles: Lucida Sans 12, 1.5 spacing
def set_font(style, size=None, bold=None):
    style.font.name = FONT
    rfonts = style.element.rPr.rFonts
    rfonts.set(qn("w:eastAsia"), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):   # theme fonts override FONT
        rfonts.attrib.pop(qn(a), None)
    ppr = style.element.pPr
    if ppr is not None and ppr.numPr is not None:   # template headings are auto-numbered; we number by hand
        ppr.remove(ppr.numPr)
        style.paragraph_format.left_indent = style.paragraph_format.first_line_indent = Cm(0)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)

normal = doc.styles["Normal"]
set_font(normal, 12)
normal.paragraph_format.line_spacing = 1.5
for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 12)]:
    set_font(doc.styles[name], size, True)
    doc.styles[name].paragraph_format.line_spacing = 1.5

code_style = doc.styles.add_style("CodeBlock", 1)
code_style.base_style = normal
code_style.font.name = "Consolas"
code_style.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
code_style.font.size = Pt(8.5)
code_style.paragraph_format.line_spacing = 1.0
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.left_indent = Cm(0.3)

# ---- page numbers in footer
for section in doc.sections:
    section.left_margin = section.right_margin = Cm(2.5)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        run._r.append(el)


def shade(par, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    par._p.get_or_add_pPr().append(shd)


def mono_block(text, fill):
    for line in text.rstrip("\n").split("\n"):
        par = doc.add_paragraph(line if line else " ", style="CodeBlock")
        shade(par, fill)
    doc.add_paragraph()


def label(text):
    par = doc.add_paragraph()
    par.add_run(text).bold = True


def add_cells(tags):
    tags = [tags] if isinstance(tags, str) else tags
    label("Code")
    mono_block("\n\n".join(CELLS[t].source for t in tags), "F2F2F2")
    label("Result")
    for t in tags:
        text = "".join(o.get("text", "") for o in CELLS[t].outputs if o.output_type == "stream")
        if text.strip():
            mono_block(text, "FFFFFF")
        for o in CELLS[t].outputs:
            if "image/png" in o.get("data", {}):
                doc.add_picture(io.BytesIO(base64.b64decode(o["data"]["image/png"])), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
for kind, *args in CONTENT:
    if kind.startswith("h"):
        doc.add_heading(args[0], level=int(kind[1]))
    elif kind == "p":
        doc.add_paragraph(args[0])
    elif kind == "bullet":
        doc.add_paragraph(args[0], style="List Paragraph").paragraph_format.left_indent = Cm(0.8)
    elif kind == "code":
        add_cells(args[0])
    elif kind == "comment":
        label("Comment on the result")
        doc.add_paragraph(args[0])
    elif kind == "pagebreak":
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

doc.save("Group5_Report.docx")
print("saved Group5_Report.docx")
